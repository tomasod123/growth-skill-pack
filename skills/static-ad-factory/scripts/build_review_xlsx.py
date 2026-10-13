#!/usr/bin/env python3
"""Build the review / run-list spreadsheet locally (xlsx with a real Approve dropdown).

Two modes:
  review   one row per rendered variant, for the first approve/reject round
           python3 build_review_xlsx.py review [--concepts data/concepts_final.json] [--renders renders] [--links links.csv]
  run      the approved set in launch order, one link column per Meta size, for whoever loads the ads
           python3 build_review_xlsx.py run [--order data/launch_order.json] [--links links.csv]

--links is an optional CSV 'filename,url' (for example exported from your file host) so each row links
to the actual image; without it the local file name is shown.

Optional publish over an existing Google Sheet (keeps its URL and the dropdowns, no Sheets API needed):
  --patch-sheet-id <SHEET_ID> --rclone-remote gdrive   (uses your rclone Google Drive OAuth token)
Gotcha: in OOXML, showDropDown=True HIDES the in-cell arrow. It is set False here on purpose.
"""
import argparse, configparser, csv, json, os, re, subprocess, sys, urllib.request
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, PatternFill, Alignment
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from adkit_common import staged_name

RATIOS = ['1x1 - 1080x1080', '4x5 - 1080x1350', '9x16 - 1080x1920']
HEAD = '333333'


def style_header(ws):
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor=HEAD)
        cell.alignment = Alignment(vertical='center', wrap_text=True)


def dropdown(ws, col, last):
    dv = DataValidation(type='list', formula1='"Approve,Reject,Maybe"', allow_blank=True)
    dv.showDropDown = False  # inverted in the spec: False SHOWS the arrow
    ws.add_data_validation(dv); dv.add(f'{col}2:{col}{last}')


def rclone_token(remote):
    subprocess.run(['rclone', 'lsd', f'{remote}:', '--max-depth', '1'], capture_output=True, timeout=90)  # forces refresh
    c = configparser.ConfigParser(); c.read(os.path.expanduser('~/.config/rclone/rclone.conf'))
    return json.loads(c[remote]['token'])['access_token']


def patch_sheet(path, sheet_id, remote):
    tok = rclone_token(remote)
    req = urllib.request.Request(f'https://www.googleapis.com/upload/drive/v3/files/{sheet_id}?uploadType=media&supportsAllDrives=true',
                                 data=open(path, 'rb').read(), method='PATCH',
                                 headers={'Authorization': f'Bearer {tok}',
                                          'Content-Type': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'})
    res = json.loads(urllib.request.urlopen(req, timeout=180).read())
    print('published over sheet:', res.get('id'), res.get('name'))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('mode', choices=['review', 'run'])
    ap.add_argument('--concepts', default='data/concepts_final.json'); ap.add_argument('--renders', default='renders')
    ap.add_argument('--order', default='data/launch_order.json'); ap.add_argument('--links')
    ap.add_argument('--out'); ap.add_argument('--patch-sheet-id'); ap.add_argument('--rclone-remote', default='gdrive')
    a = ap.parse_args()
    links = {}
    if a.links:
        for row in csv.reader(open(a.links, newline='')):
            if len(row) >= 2: links[row[0].strip()] = row[1].strip()
    wb = Workbook(); ws = wb.active
    if a.mode == 'review':
        ws.title = 'REVIEW'
        cs = {c['concept_id']: c for c in json.load(open(a.concepts))}
        hdr = ['Approve', 'Notes', 'ID', 'Image', 'Structure', 'Headline on image', 'Text elements', 'Reason to buy',
               'Needs asset', 'Facebook headline', 'Primary text', 'Watch', 'Built from']
        ws.append(hdr)
        for f in sorted(os.listdir(a.renders)) if os.path.isdir(a.renders) else []:
            m = re.match(r'([A-Z]+\d+)([ab])_.*\.png$', f)
            if not m or m.group(1) not in cs: continue
            c = cs[m.group(1)]; key = m.group(1) + m.group(2)
            ref = c.get('reference_ad_id') or ''
            ws.append(['', '', key, links.get(f, f), c.get('family_id'), c.get('headline_on_image') or '(no text)',
                       c.get('text_lines'), c.get('reason_to_buy'), c.get('needs_asset'), c.get('fb_headline'),
                       c.get('primary_text'), c.get('risk', ''),
                       f'https://www.facebook.com/ads/library/?id={ref}' if ref else ''])
        col = 'A'
    else:
        ws.title = 'READY TO RUN'
        order = json.load(open(a.order))
        hdr = ['Run order', 'Tier', 'Approve', 'ID', 'Reason to buy', 'Style', 'Headline on image', 'Text elements',
               'Why it ranks here', *RATIOS, 'Facebook headline', 'Primary text', 'Uploaded to Meta (date)']
        ws.append(hdr)
        for r in order:
            name = staged_name(r)
            ws.append([r['rank'], r['tier'], '', r['variant'], r['painpoint'], r.get('style'), r['headline'], r['text_lines'],
                       r.get('why'), *[links.get(f'{rt}/{name}', links.get(name, f'{rt}/{name}')) for rt in RATIOS],
                       r.get('fb_headline', ''), r.get('primary', ''), ''])
        col = 'C'
    style_header(ws); last = ws.max_row
    dropdown(ws, col, last); ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows(min_row=2):
        for cell in row: cell.alignment = Alignment(wrap_text=True, vertical='top')
    out = a.out or f'data/{a.mode}_sheet.xlsx'
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True); wb.save(out)
    print(f'wrote {out} | {last - 1} rows')
    if a.patch_sheet_id: patch_sheet(out, a.patch_sheet_id, a.rclone_remote)


if __name__ == '__main__':
    main()
