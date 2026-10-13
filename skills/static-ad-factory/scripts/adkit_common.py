"""Shared helpers for the static-ad-factory scripts (naming, review-file loading, flat-row maths)."""
import csv, json, os, re

FLAT = 9  # per-row max channel std below this = flat background, safe to crop


def safe(s, n=54):
    s = re.sub(r'[\\/:*?"<>|]', '', str(s or '')).strip()
    return re.sub(r'\s+', ' ', s)[:n].rstrip(' .')


def staged_name(r):
    """NN TIER PAINPOINT IDv - headline.png : sorts into launch order, identical across ratio folders."""
    return f"{r['rank']:02d} {r['tier']} {safe(r.get('painpoint'), 12)} {r['variant']} - {safe(r.get('headline'))}.png"


def load_review(path):
    """Rows from the review file: .xlsx (first sheet), .csv or .json. Needs an 'ID' (concept id + variant,
    e.g. A01a) and an 'Approve' column (Approve / Reject / Maybe)."""
    if path.endswith('.json'):
        return json.load(open(path))
    if path.endswith('.csv'):
        return list(csv.DictReader(open(path, newline='')))
    from openpyxl import load_workbook
    ws = load_workbook(path, data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(h or '').strip() for h in rows[0]]
    return [dict(zip(hdr, r)) for r in rows[1:] if any(v is not None for v in r)]


def variant_of(row):
    v = str(row.get('ID') or '').strip()
    if re.fullmatch(r'[A-Z]+\d+[ab]', v):
        return v
    m = re.search(r'\b([A-Z]+\d+[ab])\b', str(row.get('Ad file name') or row.get('file') or ''))
    return m.group(1) if m else None


def flat_rows(a, edge, limit):
    """Consecutive near-uniform rows from an edge (numpy HxWx3 array). These are safe to crop."""
    n = 0
    rows = range(a.shape[0]) if edge == 'top' else range(a.shape[0] - 1, -1, -1)
    for i in rows:
        if a[i].std(axis=0).max() > FLAT:
            break
        n += 1
        if n >= limit:
            break
    return n


def band_colour(im, edge):
    """Median colour of the outermost ~10 px strip. Solid fill beats row replication, which can
    duplicate a CTA into the band."""
    import numpy as np
    a = np.asarray(im.convert('RGB'))
    strip = {'top': a[:10], 'bottom': a[-10:], 'left': a[:, :10], 'right': a[:, -10:]}[edge]
    return tuple(int(v) for v in np.median(strip.reshape(-1, 3), axis=0))


def extend_solid(im, tw, th):
    """Grow to tw x th with a SOLID fill sampled from each edge. Never crops, never smears type."""
    from PIL import Image
    W, H = im.size; out = im
    if th > H:
        pad = th - H; t, b = pad // 2, pad - pad // 2
        c = Image.new('RGB', (W, th))
        if t: c.paste(Image.new('RGB', (W, t), band_colour(out, 'top')), (0, 0))
        c.paste(out, (0, t))
        if b: c.paste(Image.new('RGB', (W, b), band_colour(out, 'bottom')), (0, t + H))
        out = c
    W, H = out.size
    if tw > W:
        pad = tw - W; l, r = pad // 2, pad - pad // 2
        c = Image.new('RGB', (tw, H))
        if l: c.paste(Image.new('RGB', (l, H), band_colour(out, 'left')), (0, 0))
        c.paste(out, (l, 0))
        if r: c.paste(Image.new('RGB', (r, H), band_colour(out, 'right')), (l + W, 0))
        out = c
    return out


def find_render(folder, variant):
    """renders/<variant>_<family>.png. Filters iterdir, not glob: names with [] break glob."""
    if not os.path.isdir(folder):
        return None
    for f in sorted(os.listdir(folder)):
        if f.startswith(variant + '_') and f.endswith('.png'):
            return os.path.join(folder, f)
    return None
